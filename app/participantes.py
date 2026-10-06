"""Leitura tolerante de participantes a partir de um arquivo .docx."""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from docx import Document

SEPARADORES = re.compile(r"\s+[—–-]\s+|\s*[:|;(]\s*|\s*\t\s*|\s{3,}")
MARCADORES = re.compile(r"^\s*(?:[-•*·▪●◦–—]+|\d+[.)]|[a-zA-Z][.)])\s*")
CABECALHO_PART = re.compile(r"^\s*(participantes?|entrevistad[oa]s?|presentes|pessoas|integrantes)\b\s*:?\s*(.*)$", re.I)
CABECALHO_FIM = re.compile(r"^\s*(objetivo|pauta|roteiro|perguntas|t[oó]picos|agenda|observa[cç][oõ]es|resumo|contexto|data|local|dura[cç][aã]o)\b", re.I)
ID_ENTREVISTA = re.compile(r"\b(?:entrevista|interview)\b\s*(?:n[º°o.]*\s*)?[:#-]?\s*([A-Z0-9][\w./-]*\d[\w./-]*|[A-Z]{2,}[\w./-]*)", re.I)
CODIGO = re.compile(r"\b[A-Z]{2,}(?:-[A-Z0-9]+)+\b")
PALAVRAS_NAO_NOME = {
    "entrevista", "participantes", "participante", "transcrição", "transcricao", "data", "local",
    "horário", "horario", "objetivo", "pauta", "duração", "duracao", "projeto", "empresa", "nome",
    "cargo", "função", "funcao", "área", "area", "papel", "observações", "observacoes",
}
CONECTORES = {"da", "de", "do", "das", "dos", "e", "di", "du", "del", "van", "von", "y"}


def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def parece_nome(texto: str) -> bool:
    t = texto.strip().strip(".,;")
    if not (2 <= len(t) <= 60) or re.search(r"[\d@/\\#=<>{}\[\]]", t):
        return False
    palavras = t.split()
    if not 1 <= len(palavras) <= 6:
        return False
    if _sem_acento(palavras[0].lower()) in {_sem_acento(p) for p in PALAVRAS_NAO_NOME}:
        return False
    significativas = [p for p in palavras if p.lower() not in CONECTORES]
    if not significativas:
        return False
    return all(p[0].isupper() for p in significativas if p[0].isalpha())


def _dividir(linha: str) -> tuple[str, str]:
    linha = MARCADORES.sub("", linha).strip()
    partes = [p.strip(" )") for p in SEPARADORES.split(linha, maxsplit=1) if p and p.strip(" )")]
    if not partes:
        return "", ""
    nome = partes[0]
    papel = partes[1] if len(partes) > 1 else ""
    return nome, papel


def _linhas(doc: Document) -> list[str]:
    linhas: list[str] = []
    for p in doc.paragraphs:
        for sub in p.text.splitlines():
            if sub.strip():
                linhas.append(sub.strip())
    return linhas


def _tabelas(doc: Document) -> list[dict]:
    achados: list[dict] = []
    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = []
            for c in linha.cells:
                t = c.text.strip()
                if not celulas or celulas[-1] != t:  # células mescladas repetem o texto
                    celulas.append(t)
            celulas = [c for c in celulas if c]
            if not celulas:
                continue
            idx = next((i for i, c in enumerate(celulas) if parece_nome(c)), None)
            if idx is None:
                continue
            papel = " / ".join(c for i, c in enumerate(celulas) if i != idx and len(c) < 60)
            achados.append({"nome": celulas[idx].strip(), "papel": papel})
    return achados


def extrair(caminho: str | Path) -> dict:
    """Retorna {'entrevista': str, 'participantes': [{'nome','papel'}], 'texto': str}."""
    doc = Document(str(caminho))
    linhas = _linhas(doc)
    texto = "\n".join(linhas)

    entrevista = ""
    for linha in linhas[:15]:
        m = ID_ENTREVISTA.search(linha)
        if m and not re.match(r"^(com|de|do|da)$", m.group(1), re.I):
            entrevista = m.group(1).strip(" .:-")
            break
    if not entrevista:
        m = CODIGO.search(texto)
        entrevista = m.group(0) if m else ""

    candidatos: list[dict] = []
    # 1) Bloco iniciado por "Participantes:" (ou similar)
    dentro = False
    for linha in linhas:
        cab = CABECALHO_PART.match(linha)
        if cab:
            dentro = True
            resto = cab.group(2).strip()
            if resto:  # "Participantes: Fulano, Beltrano e Sicrano"
                for pedaco in re.split(r",|;|\s+e\s+", resto):
                    nome, papel = _dividir(pedaco)
                    if parece_nome(nome):
                        candidatos.append({"nome": nome, "papel": papel})
            continue
        if dentro:
            if CABECALHO_FIM.match(linha):
                dentro = False
                continue
            nome, papel = _dividir(linha)
            if parece_nome(nome):
                candidatos.append({"nome": nome, "papel": papel})
            elif candidatos and len(linha) > 80:
                dentro = False  # chegou em texto corrido

    # 2) Tabelas
    if not candidatos:
        candidatos = _tabelas(doc)

    # 3) Linhas no formato "Nome — Papel" em qualquer lugar do documento
    if not candidatos:
        for linha in linhas:
            if re.search(r"\s[—–-]\s|:\s", linha) or MARCADORES.match(linha):
                nome, papel = _dividir(linha)
                if parece_nome(nome) and papel:
                    candidatos.append({"nome": nome, "papel": papel})

    vistos: set[str] = set()
    participantes = []
    for c in candidatos:
        chave = _sem_acento(c["nome"].lower())
        if chave not in vistos:
            vistos.add(chave)
            participantes.append({"nome": c["nome"].strip(), "papel": c["papel"].strip()})
    return {"entrevista": entrevista, "participantes": participantes, "texto": texto[:5000]}
