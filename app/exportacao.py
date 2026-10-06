"""Exportação para DOCX, XLSX, TXT, JSON e SRT."""
from __future__ import annotations

import io
import json

from . import tempo
from .projeto import nome_exibicao


def _participantes(proj: dict) -> list[dict]:
    return [p for p in proj.get("participantes", []) if p.get("nome")]


def _linhas(proj: dict) -> list[dict]:
    return [{"inicio": s["inicio"], "fim": s["fim"], "pessoa": nome_exibicao(proj, s),
             "speaker": s.get("speaker"), "texto": s["texto"]} for s in proj.get("segmentos", [])]


def para_json(proj: dict) -> bytes:
    dados = {
        "entrevista": proj.get("entrevista", ""),
        "arquivo_video": proj.get("arquivo_video"),
        "criado_em": proj.get("criado_em"),
        "atualizado_em": proj.get("atualizado_em"),
        "idioma": proj.get("config", {}).get("idioma") or proj.get("info", {}).get("idioma_detectado"),
        "modelo_transcricao": proj.get("config", {}).get("modelo"),
        "duracao_segundos": proj.get("info", {}).get("duracao"),
        "participantes": _participantes(proj),
        "speakers": proj.get("speakers", {}),
        "segmentos": [
            {"id": i, "inicio": s["inicio"], "fim": s["fim"],
             "inicio_fmt": tempo.formatar(s["inicio"]), "fim_fmt": tempo.formatar(s["fim"]),
             "speaker": s["speaker"], "pessoa": s["pessoa"], "texto": s["texto"]}
            for i, s in enumerate(_linhas(proj), 1)
        ],
    }
    return json.dumps(dados, ensure_ascii=False, indent=2).encode("utf-8")


def para_txt(proj: dict) -> bytes:
    linhas = ["TRANSCRIÇÃO DA ENTREVISTA", ""]
    if proj.get("entrevista"):
        linhas.append(f"Entrevista: {proj['entrevista']}")
    parts = _participantes(proj)
    if parts:
        linhas.append("Participantes:")
        linhas += [f"- {p['nome']}" + (f" — {p['papel']}" if p.get("papel") else "") for p in parts]
    linhas += ["", "-" * 48, ""]
    for s in _linhas(proj):
        linhas.append(f"[{tempo.formatar(s['inicio'])} - {tempo.formatar(s['fim'])}] {s['pessoa'].upper()}")
        linhas.append(s["texto"])
        linhas.append("")
    return "\n".join(linhas).encode("utf-8")


def para_srt(proj: dict) -> bytes:
    blocos = []
    for i, s in enumerate(_linhas(proj), 1):
        blocos.append(f"{i}\n{tempo.formatar_srt(s['inicio'])} --> {tempo.formatar_srt(s['fim'])}\n"
                      f"{s['pessoa']}: {s['texto']}\n")
    return "\n".join(blocos).encode("utf-8")


def para_docx(proj: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run("TRANSCRIÇÃO DA ENTREVISTA")
    r.bold = True
    r.font.size = Pt(16)

    if proj.get("entrevista"):
        p = doc.add_paragraph()
        p.add_run("Entrevista: ").bold = True
        p.add_run(proj["entrevista"])
    parts = _participantes(proj)
    if parts:
        doc.add_paragraph().add_run("Participantes:").bold = True
        for pt in parts:
            doc.add_paragraph(pt["nome"] + (f" — {pt['papel']}" if pt.get("papel") else ""),
                              style="List Bullet")
    doc.add_paragraph("─" * 48)

    for s in _linhas(proj):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(0)
        rt = p.add_run(f"[{tempo.formatar(s['inicio'])} - {tempo.formatar(s['fim'])}]")
        rt.font.size = Pt(9)
        rt.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.add_run(s["pessoa"].upper()).bold = True
        doc.add_paragraph(s["texto"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def para_xlsx(proj: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Transcrição"
    cab = ["Início", "Fim", "Pessoa", "Texto", "Speaker (diarização)"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5496")
    for s in _linhas(proj):
        ws.append([tempo.formatar(s["inicio"]), tempo.formatar(s["fim"]), s["pessoa"], s["texto"], s["speaker"]])
    for col, larg in zip("ABCDE", (11, 11, 24, 100, 20)):
        ws.column_dimensions[col].width = larg
    for linha in ws.iter_rows(min_row=2):
        linha[3].alignment = Alignment(wrap_text=True, vertical="top")
        for c in (linha[0], linha[1], linha[2], linha[4]):
            c.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wp = wb.create_sheet("Participantes")
    wp.append(["Entrevista", proj.get("entrevista", "")])
    wp.append([])
    wp.append(["Nome", "Papel"])
    for c in wp[3]:
        c.font = Font(bold=True)
    for pt in _participantes(proj):
        wp.append([pt["nome"], pt.get("papel", "")])
    wp.column_dimensions["A"].width = 30
    wp.column_dimensions["B"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FORMATOS = {
    "docx": (para_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "xlsx": (para_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "txt": (para_txt, "text/plain; charset=utf-8"),
    "json": (para_json, "application/json"),
    "srt": (para_srt, "application/x-subrip"),
}
