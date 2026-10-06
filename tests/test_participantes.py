from docx import Document

from app.participantes import extrair, parece_nome


def _docx(tmp_path, linhas, tabela=None):
    doc = Document()
    for l in linhas:
        doc.add_paragraph(l)
    if tabela:
        t = doc.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, v in enumerate(linha):
                t.cell(i, j).text = v
    caminho = tmp_path / "p.docx"
    doc.save(caminho)
    return caminho


def test_exemplo_do_enunciado(tmp_path):
    arq = _docx(tmp_path, ["ENTREVISTA ENT-TI-03", "", "Participantes:", "",
                           "Adriano Petrovich — Entrevistador", "José Lucas — TI", "Maria Silva — Operações"])
    r = extrair(arq)
    assert r["entrevista"] == "ENT-TI-03"
    assert [p["nome"] for p in r["participantes"]] == ["Adriano Petrovich", "José Lucas", "Maria Silva"]
    assert r["participantes"][0]["papel"] == "Entrevistador"


def test_lista_com_marcadores_e_fim_de_bloco(tmp_path):
    arq = _docx(tmp_path, ["Entrevista: ENT-RH-01", "Participantes", "- Ana Paula de Souza (RH)",
                           "• Carlos Eduardo: Gerente", "Objetivo", "Entender o processo de contratação."])
    r = extrair(arq)
    assert r["entrevista"] == "ENT-RH-01"
    assert [p["nome"] for p in r["participantes"]] == ["Ana Paula de Souza", "Carlos Eduardo"]
    assert r["participantes"][0]["papel"] == "RH"


def test_participantes_na_mesma_linha(tmp_path):
    arq = _docx(tmp_path, ["Participantes: João Silva, Maria Souza e Pedro Alves"])
    assert [p["nome"] for p in extrair(arq)["participantes"]] == ["João Silva", "Maria Souza", "Pedro Alves"]


def test_tabela(tmp_path):
    arq = _docx(tmp_path, ["Reunião de levantamento"],
                tabela=[["Nome", "Área"], ["Fernanda Lima", "Financeiro"], ["Rafael Costa", "TI"]])
    r = extrair(arq)
    assert [p["nome"] for p in r["participantes"]] == ["Fernanda Lima", "Rafael Costa"]
    assert r["participantes"][1]["papel"] == "TI"


def test_sem_participantes(tmp_path):
    arq = _docx(tmp_path, ["Este documento descreve o roteiro da conversa em texto corrido sem nomes."])
    assert extrair(arq)["participantes"] == []


def test_parece_nome():
    assert parece_nome("Maria da Silva")
    assert not parece_nome("Data: 12/03/2025")
    assert not parece_nome("entrevista com o time")
