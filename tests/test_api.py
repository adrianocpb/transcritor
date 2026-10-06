"""Fluxo completo pela API, com transcrição e diarização simuladas (sem baixar modelos)."""
import io
import json
import shutil
import subprocess
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import diarizacao, transcricao
from app.main import app

pytestmark = pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="FFmpeg necessário")


@pytest.fixture
def video(tmp_path):
    arq = tmp_path / "entrevista.mp4"
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=8",
                    "-f", "lavfi", "-i", "color=c=black:s=160x120:d=8", "-shortest",
                    "-c:v", "libx264", "-c:a", "aac", str(arq)], check=True)
    return arq


@pytest.fixture
def docx_participantes():
    doc = Document()
    for l in ["ENTREVISTA ENT-TI-03", "Participantes:", "Adriano Petrovich — Entrevistador", "José Lucas — TI"]:
        doc.add_paragraph(l)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def falso_transcrever(audio, modelo, idioma, dispositivo, threads, prompt, progresso):
    assert abs(len(audio) - 8 * 16000) < 1600 and idioma == "pt" and "Adriano Petrovich" in prompt
    progresso(0.5, "Transcrevendo…")
    pal = [(0.5, 1.0, "Bom"), (1.1, 1.6, "dia."), (1.7, 2.5, "Como"), (2.6, 3.0, "funciona?"),
           (4.0, 4.5, "Hoje"), (4.6, 5.2, "começa"), (5.3, 6.0, "assim.")]
    return ([{"inicio": a, "fim": b, "texto": " " + t, "prob": 0.95} for a, b, t in pal],
            {"idioma_detectado": "pt", "prob_idioma": 1.0, "duracao": 8.0})


def falso_diarizar(audio, dispositivo, num, mn, mx, progresso):
    assert num == 2
    return [{"inicio": 0.4, "fim": 3.2, "speaker": "SPEAKER_00"}, {"inicio": 3.9, "fim": 6.2, "speaker": "SPEAKER_01"}]


def test_fluxo_completo(monkeypatch, video, docx_participantes):
    monkeypatch.setattr(transcricao, "transcrever", falso_transcrever)
    monkeypatch.setattr(diarizacao, "diarizar", falso_diarizar)
    c = TestClient(app)

    assert "hardware" in c.get("/api/sistema").json()

    with open(video, "rb") as f:
        r = c.post("/api/projetos", files={"video": ("entrevista.mp4", f, "video/mp4"),
                                           "docx": ("p.docx", docx_participantes, "application/octet-stream")})
    assert r.status_code == 200, r.text
    proj = r.json()
    pid = proj["id"]
    assert proj["entrevista"] == "ENT-TI-03"
    assert [p["nome"] for p in proj["participantes"]] == ["Adriano Petrovich", "José Lucas"]

    r = c.put(f"/api/projetos/{pid}/participantes", json={"entrevista": "ENT-TI-03", "participantes": proj["participantes"]})
    assert r.status_code == 200
    r = c.post(f"/api/projetos/{pid}/processar", json={"modelo": "small", "idioma": "pt", "num_speakers": 2})
    assert r.status_code == 200, r.text

    for _ in range(100):
        proj = c.get(f"/api/projetos/{pid}").json()
        if proj["status"] in ("mapear_speakers", "erro"):
            break
        time.sleep(0.1)
    assert proj["status"] == "mapear_speakers", proj.get("erro")
    assert proj["speakers"] == {"SPEAKER_00": "", "SPEAKER_01": ""}
    assert [s["texto"] for s in proj["segmentos"]] == ["Bom dia. Como funciona?", "Hoje começa assim."]

    r = c.put(f"/api/projetos/{pid}/speakers", json={"speakers": {"SPEAKER_00": "Adriano Petrovich", "SPEAKER_01": "José Lucas"}})
    assert r.json()["status"] == "revisao"

    segs = proj["segmentos"]
    segs[1]["texto"] = "Hoje o processo começa assim."
    r = c.put(f"/api/projetos/{pid}/segmentos", json={"segmentos": segs})
    assert r.status_code == 200

    # player: vídeo com suporte a Range
    r = c.get(f"/api/projetos/{pid}/midia", headers={"Range": "bytes=0-99"})
    assert r.status_code == 206 and len(r.content) == 100
    assert c.get(f"/api/projetos/{pid}/midia?tipo=audio").status_code == 200

    txt = c.get(f"/api/projetos/{pid}/exportar/txt").text
    assert "[00:00:00 - 00:00:03] ADRIANO PETROVICH" in txt and "JOSÉ LUCAS" in txt

    dados = json.loads(c.get(f"/api/projetos/{pid}/exportar/json").content)
    assert dados["entrevista"] == "ENT-TI-03"
    assert dados["segmentos"][1]["pessoa"] == "José Lucas"
    assert dados["segmentos"][1]["texto"] == "Hoje o processo começa assim."

    r = c.get(f"/api/projetos/{pid}/exportar/docx")
    assert r.status_code == 200 and "Transcricao_ENT-TI-03.docx" in r.headers["content-disposition"]
    textos = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert "TRANSCRIÇÃO DA ENTREVISTA" in textos and "ADRIANO PETROVICH" in textos and "JOSÉ LUCAS" in textos

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get(f"/api/projetos/{pid}/exportar/xlsx").content))
    linhas = list(wb["Transcrição"].iter_rows(values_only=True))
    assert linhas[0][:4] == ("Início", "Fim", "Pessoa", "Texto")
    assert linhas[2][2] == "José Lucas"

    assert "-->" in c.get(f"/api/projetos/{pid}/exportar/srt").text
    assert any(p["id"] == pid for p in c.get("/api/projetos").json())
    assert c.delete(f"/api/projetos/{pid}").status_code == 200
    assert c.get(f"/api/projetos/{pid}").status_code == 404


def test_erro_no_processamento_e_reportado(monkeypatch, video):
    def quebra(*a, **k):
        raise RuntimeError("modelo ausente")
    monkeypatch.setattr(transcricao, "transcrever", quebra)
    c = TestClient(app)
    with open(video, "rb") as f:
        pid = c.post("/api/projetos", files={"video": ("v.mp4", f, "video/mp4")}).json()["id"]
    c.post(f"/api/projetos/{pid}/processar", json={"modelo": "tiny", "diarizar": False})
    for _ in range(100):
        proj = c.get(f"/api/projetos/{pid}").json()
        if proj["status"] == "erro":
            break
        time.sleep(0.1)
    assert proj["status"] == "erro" and "modelo ausente" in proj["erro"]


def test_rejeita_formato_e_id_invalidos():
    c = TestClient(app)
    assert c.post("/api/projetos", files={"video": ("x.exe", b"123", "application/octet-stream")}).status_code == 400
    assert c.get("/api/projetos/..%2F..%2Fetc").status_code == 404
